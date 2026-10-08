"""Independent small-table task contracts, executed through the existing harness.

Expected numbers are literals or hand-derived formulas, never calculated by the
production functions under test. No online service, model or user input is used.
"""
from __future__ import annotations

from copy import deepcopy
from dataclasses import replace
from io import BytesIO
import json
import math
from unittest.mock import patch, Mock
from zipfile import ZipFile

import pandas as pd
from scipy import stats

from insightpilot.evaluation.models import EvaluationCase, EvaluationSuiteResult
from insightpilot.evaluation.harness import EvaluationHarness


FAMILIES = {
    "revenue": "订单口径与净/毛收入", "conversion": "转化率澄清", "users": "去重人数与行数",
    "weighted": "加权比率与均值比率", "branches": "维度分支与窗口对比", "comparability": "不可比单位/口径",
    "fanout": "连接放大与跨事实边界", "sql": "安全作用域与危险分支", "experiment": "实验单位/方向/区间",
    "invalidation": "缓存失效与历史不可变", "release": "释放与分支归属", "exports": "六类导出绑定",
}


def daily_fixture() -> pd.DataFrame:
    """16 rows: previous day orders=30/revenue=300; current=20/200."""
    rows = []
    for day in range(8):
        for city, channel, before, after, visits, attempts in (("甲", "搜索", 10, 8, 20, 15), ("乙", "直接", 20, 12, 100, 30)):
            orders = before if day < 7 else after
            rows.append({"date": pd.Timestamp("2026-01-01") + pd.Timedelta(days=day), "city": city, "channel": channel,
                "orders": orders, "revenue": orders * 10, "visitors": visits, "payment_attempts": attempts,
                "successful_payments": orders, "cvr": orders / visits, "payment_success_rate": orders / attempts})
    return pd.DataFrame(rows)


def _metadata(frame: pd.DataFrame, name="daily") -> dict:
    # Guided preflight requires real exact value facts; bare column names do not
    # prove numeric/date readiness. This changes fixture preparation, not oracle.
    from insightpilot.agents.dataset import PreparedDataset
    return PreparedDataset.from_tables({name: frame}, dataset_id="task-fixture", revision="1").metadata


def _mapping(metric="revenue", dimension="city") -> dict:
    return {"table_name": "daily", "date_column": "date", "metric_columns": [metric],
        "dimension_columns": [dimension], "aggregation": "sum", "unit": "元" if metric == "revenue" else "比例" if metric == "cvr" else "单"}


def _parameters(current="2026-01-08") -> dict:
    return {"current_start": current, "current_end": current, "previous_start": "2026-01-07", "previous_end": "2026-01-07", "aggregation": "sum"}


def _run_period(*, metric="revenue", dimension="city", current="2026-01-08", frame=None) -> dict:
    from insightpilot.agents.workflow import run_agent_analysis
    return run_agent_analysis("按明确口径核对指标", {"daily": daily_fixture() if frame is None else frame},
        data_source_type="csv", goal_mode="metric_diagnosis", playbook_id="period_comparison",
        column_mapping=_mapping(metric, dimension), playbook_parameters=_parameters(current),
        clarification_answers={"aggregation": "sum"}, presentation_mode="deferred")


def _summary_value(result):
    frame = result["result_tables"]["period_summary"]
    return float(frame.loc[frame["comparison_period"] == "current", "metric_value"].iloc[0])


def _observed(checks: dict[str, bool], **observed) -> dict:
    checks = {str(k): bool(v) for k, v in checks.items()}
    return {"scores": {"task_contract": float(all(checks.values()))}, "checks": checks,
        "failures": [name for name, passed in checks.items() if not passed], "observed": observed}


def _close(actual, expected, tolerance=1e-10):
    return isinstance(actual, (int, float)) and math.isfinite(actual) and math.isclose(actual, expected, rel_tol=tolerance, abs_tol=tolerance)


def build_task_cases() -> list[EvaluationCase]:
    """Build fresh evaluation-local fixtures; no process-global results or cache."""
    from insightpilot.planning.planner import prepare_analysis_request
    from insightpilot.agents.workflow import run_agent_analysis
    from insightpilot.analysis.metric_diagnosis import _daily_series
    from insightpilot.analysis.threads import AnalysisThread, compare_runs
    from insightpilot.analysis.experiments import analyze_ab_test
    from insightpilot.tools.duckdb_engine import AnalyticsEngine
    from insightpilot.tools.sql_safety import validate_sql
    from app.session_data import AnalysisSession, analysis_cache_key
    from insightpilot.semantic.catalog import load_builtin_catalog
    from insightpilot.semantic.compiler import MetricCompiler, execute_query_plan
    from insightpilot.semantic.query import MetricRequest

    frame = daily_fixture()
    schema = _metadata(frame)
    cases: list[EvaluationCase] = []
    outcomes: dict[tuple, dict] = {}
    def result(metric="revenue", dimension="city", current="2026-01-08"):
        key = (metric, dimension, current)
        if key not in outcomes:
            outcomes[key] = _run_period(metric=metric, dimension=dimension, current=current)
        return outcomes[key]
    def case(case_id, family, question, runner, expected, *, fixture="daily16", evidence=("实际执行状态",), forbidden=("把未执行解释为数值0",)):
        cases.append(EvaluationCase(case_id, "task", question, runner, family=FAMILIES[family], question=question,
            input_fixture={"name": fixture, "source": "evaluation.task_suite中的固定小表；数字预期在case中独立声明"},
            expected=expected, required_evidence=evidence, forbidden_behaviors=forbidden))
    def blocked(question, expected_status):
        with patch.object(AnalyticsEngine, "run_parameterized_sql", side_effect=AssertionError("阻塞时执行了SQL")) as query:
            actual = run_agent_analysis(question, {"daily": frame}, data_source_type="csv", presentation_mode="deferred")
        return _observed({"规划状态符合预期": actual.get("planning_status") == expected_status,
            "阻塞状态没有查询": query.call_count == 0, "未生成已完成的数值结果": actual.get("execution_status") != "COMPLETED"},
            planning_status=actual.get("planning_status"), query_count=query.call_count, clarification=actual.get("clarification"))
    def gross():
        actual = result()
        card = actual.get("metric_definitions", [{}])[0]
        return _observed({"流程完成": actual.get("execution_status") == "COMPLETED", "支付金额为200": _close(_summary_value(actual), 200),
            "金额口径未冒充净收入": card.get("metric_id") == "revenue" and card.get("expression") == "sum(revenue)" and "不自动认定净收入" in str(card.get("business_meaning")),
            "金额单位为元": card.get("unit") == "元"}, actual=_summary_value(actual), metric=card)
    case("revenue_gross_200", "revenue", "明确支付金额（不扣退款）", gross, {"status": "COMPLETED", "metric": "revenue", "value": 200, "unit": "元"})
    case("revenue_ambiguous", "revenue", "收入是多少？", lambda: blocked("收入是多少？", "needs_clarification"), {"planning_status": "needs_clarification", "queries": 0})
    case("revenue_net_unsupported", "revenue", "扣退款净收入是多少？", lambda: blocked("扣退款净收入是多少？", "unsupported"), {"planning_status": "unsupported", "queries": 0})
    case("conversion_ambiguous", "conversion", "转化率为什么下降？", lambda: blocked("转化率为什么下降？", "needs_clarification"), {"planning_status": "needs_clarification", "queries": 0})
    def selected_conversion():
        prepared = prepare_analysis_request("转化率为什么下降？", table_metadata=schema, playbook_id="period_comparison",
            column_mapping={"table_name": "daily", "date_column": "date", "dimension_columns": ["city"], "aggregation": "sum"},
            playbook_parameters=_parameters(), clarification_answers={"metric": "daily:cvr"})
        request = prepared["normalized_request"]
        actual = run_agent_analysis(**request, tables={"daily": frame}, data_source_type="csv", presentation_mode="deferred", clarification_answers={"metric": "daily:cvr"})
        card = actual["metric_definitions"][0]
        return _observed({"澄清后可执行": prepared["planning_status"] == "ready" and actual["execution_status"] == "COMPLETED",
            "采用访问到订单口径": (card.get("numerator"), card.get("denominator")) == ("orders", "visitors"),
            "当前转化率20除120": _close(_summary_value(actual), 1 / 6)}, actual=_summary_value(actual), numerator=card.get("numerator"), denominator=card.get("denominator"))
    case("conversion_selected_weighted", "conversion", "选择访问到下单转化率后执行", selected_conversion, {"value": 1 / 6, "numerator": "orders", "denominator": "visitors", "tolerance": 1e-10})
    people = pd.DataFrame({"date": ["2026-01-01"] * 5, "user_id": [1, 1, 2, 2, 3]})
    def distinct():
        actual = float(_daily_series(people, "active_users")["metric_value"].iloc[0])
        return _observed({"按实体去重为3": actual == 3, "不是5行": actual != 5}, actual=actual)
    case("users_distinct_three", "users", "明确按user_id去重统计用户", distinct, {"distinct_users": 3, "rows": 5}, fixture="people5:user_id=[1,1,2,2,3]")
    def row_count():
        with AnalyticsEngine() as engine:
            engine.register_tables({"people": people})
            actual = engine.run_sql('SELECT COUNT(*) AS rows, COUNT(DISTINCT user_id) AS users FROM people').iloc[0].to_dict()
        return _observed({"行数为5": actual["rows"] == 5, "用户为3": actual["users"] == 3}, actual=actual)
    case("users_rows_are_not_uv", "users", "分别统计观测行与去重人数", row_count, {"rows": 5, "users": 3}, fixture="people5")
    ratios = pd.DataFrame({"date": ["2026-01-01"] * 2, "clicks": [1, 80], "impressions": [10, 100], "ctr": [.1, .8]})
    def weighted():
        actual = float(_daily_series(ratios, "ctr")["metric_value"].iloc[0])
        return _observed({"按81除110聚合": _close(actual, 81 / 110), "不是均值0.45": not _close(actual, .45)}, actual=actual)
    case("weighted_ratio_not_mean", "weighted", "两个样本量不同分组的总体点击率", weighted, {"weighted": 81 / 110, "forbidden_mean": .45}, fixture="ratio2:[1/10,80/100]")
    def zero_ratio():
        actual = _daily_series(ratios.assign(impressions=0), "ctr")["metric_value"].iloc[0]
        return _observed({"零分母未定义": pd.isna(actual)}, is_undefined=bool(pd.isna(actual)))
    case("weighted_zero_denominator", "weighted", "总曝光为0时点击率不可计算", zero_ratio, {"undefined": True}, fixture="ratio2:denominator=0")
    def pair(metric="revenue", right_dimension="city", right_current="2026-01-08"):
        thread = AnalysisThread()
        first = result(metric)
        second = result(metric, right_dimension, right_current)
        # Distinct actual runs are necessary even for equal configurations.
        if first is second:
            second = _run_period(metric=metric)
        a = thread.add(first, {"column_mapping": _mapping(metric), "parameters": _parameters()}, dataset_id="fixture", revision="1", schema={"daily": list(frame)})
        b = thread.add(second, {"column_mapping": _mapping(metric, right_dimension), "parameters": _parameters(right_current)}, dataset_id="fixture", revision="1", schema={"daily": list(frame)})
        return thread, a, b
    def dimensions():
        _, a, b = pair(right_dimension="channel")
        comparison = compare_runs(a, b)
        return _observed({"维度不同拒绝逐行相减": comparison["status"] == "not_comparable" and not comparison["rows"],
            "标明维度原因": any("维度" in x for x in comparison["reasons"])}, comparison=comparison)
    case("branch_dimension_diff", "branches", "按城市与按渠道的两个分支比较", dimensions, {"status": "not_comparable", "numeric_rows": 0}, evidence=("配置差异", "原运行摘要"))
    def windows():
        _, a, b = pair(right_current="2026-01-07")
        comparison = compare_runs(a, b)
        row = comparison["rows"][0] if comparison["rows"] else {}
        return _observed({"窗口差异标记": comparison["status"] == "context_differs", "金额差100": _close(row.get("absolute_change"), 100),
            "相对差0.5": _close(row.get("relative_change"), .5), "只作描述性差异": "不是显著性检验" in row.get("note", "")}, comparison=comparison)
    case("branch_window_descriptive", "branches", "同一支付金额口径的两个日期窗口", windows, {"status": "context_differs", "run_a": 200, "run_b": 300, "absolute": 100, "relative": .5})
    def incomparable(field, value):
        _, a, b = pair()
        context = b.context
        context["definitions"][0][field] = value
        b = replace(b, comparison_context=json.dumps(context, ensure_ascii=False))
        actual = compare_runs(a, b)
        return _observed({"口径不同拒绝": actual["status"] == "not_comparable", "没有数字差异": actual["rows"] == []}, comparison=actual)
    case("comparison_unit_rejected", "comparability", "金额元与美元不可直接比较", lambda: incomparable("unit", "USD"), {"status": "not_comparable"})
    case("comparison_definition_rejected", "comparability", "净收入定义不可与支付金额相减", lambda: incomparable("expression", "sum(revenue)-sum(refund_amount)"), {"status": "not_comparable"})
    def fanout():
        compiler = MetricCompiler(load_builtin_catalog().get("commerce_demo"))
        plan = compiler.compile(MetricRequest(metrics=["total_revenue"], dimensions=["product_category"]))
        with patch.object(AnalyticsEngine, "run_parameterized_sql", side_effect=AssertionError("未批准fanout执行")) as query:
            output, review = execute_query_plan(plan, {})
        return _observed({"fanout需审批或不可执行": plan.requires_approval or not plan.executable,
            "未审批不执行": output is None and query.call_count == 0}, risk=plan.risk_level, decision=review.decision, query_count=query.call_count)
    case("fanout_no_silent_sum", "fanout", "订单金额按商品分类关联明细时阻止重复求和", fanout, {"execute": False, "queries": 0}, fixture="commerce semantic relationship:orders->items one_to_many")
    def cross_fact():
        try:
            MetricCompiler(load_builtin_catalog().get("commerce_demo")).compile(MetricRequest(metrics=["revenue_per_session"], dimensions=["order_date"], date_dimension="order_date"))
        except ValueError as exc:
            return _observed({"明确跨事实时间不支持": "跨事实" in str(exc)}, reason=str(exc))
        return _observed({"拒绝未定义跨事实时间": False})
    case("cross_fact_time_unsupported", "fanout", "跨订单与会话事实按订单日期求比率", cross_fact, {"rejected": True}, fixture="commerce semantic model")
    def safe_cte():
        sql = 'WITH c AS (SELECT "城市", "金额", "备注" FROM t WHERE "金额" >= ?) SELECT "城市", SUM("金额") AS total FROM c WHERE "备注" = ? GROUP BY "城市"'
        table = pd.DataFrame({"城市": ["甲", "甲", "乙"], "金额": [10, 20, 30], "备注": ["DROP TABLE", "DROP TABLE", "normal"]})
        validation = validate_sql(sql, allowed_tables={"t"}, table_columns={"t": list(table)}, dialect="duckdb")
        with AnalyticsEngine() as engine:
            engine.register_tables({"t": table}); actual = engine.run_parameterized_sql(sql, [10, "DROP TABLE"])
        return _observed({"参数化CTE允许": validation.allowed, "只引用真实授权表": validation.referenced_tables == ("t",),
            "字符串写词不误伤且甲金额30": actual.to_dict("records") == [{"城市": "甲", "total": 30}]}, validation=validation.to_dict(), actual=actual.to_dict("records"))
    case("sql_safe_cte_chinese_bound", "sql", "中文列名CTE与含DROP文本的参数化聚合", safe_cte, {"rows": [{"城市": "甲", "total": 30}], "physical_tables": ["t"]}, fixture="sql3:10+20=30")
    def unsafe_union():
        sql = 'SELECT x FROM t UNION ALL SELECT x FROM forbidden'
        with AnalyticsEngine() as engine:
            engine.register_tables({"t": pd.DataFrame({"x": [1]})})
            connection = engine._connection
            proxy = Mock(wraps=connection); engine._connection = proxy
            try: engine.run_sql(sql); rejected = False
            except ValueError: rejected = True
            calls = proxy.execute.call_count
            engine._connection = connection
        return _observed({"危险后续分支拒绝": rejected, "用户SQL零执行": calls == 0}, execute_count=calls)
    case("sql_union_forbidden_zero_execution", "sql", "UNION后续分支访问未授权表", unsafe_union, {"rejected": True, "execute_count": 0}, fixture="t:x=[1]", forbidden=("执行任何用户SQL", "忽略UNION后续分支"))
    experiment = pd.DataFrame({"user_id": [1, 1, 2, 2, 3, 3, 4, 4], "group": ["control"] * 4 + ["treatment"] * 4, "outcome": [1, 3, 3, 5, 3, 5, 5, 7]})
    def experiment_units():
        actual = analyze_ab_test(experiment, "group", "outcome", statistical_unit="user_id")
        # Hand derivation: user means [2,4]/[4,6], each s²=2, SE=sqrt(2), Welch df=2.
        p_expected = 1 - 1 / math.sqrt(2)
        half_width = float(stats.t.ppf(.975, 2)) * math.sqrt(2)
        return _observed({"每组两个用户": actual["sample_size"] == {"control": 2, "treatment": 2},
            "差为2": _close(actual["absolute_lift"], 2), "独立闭式p值": _close(actual["p_value"], p_expected),
            "区间同口径": all(_close(x, y) for x, y in zip(actual["confidence_interval"], [2 - half_width, 2 + half_width])),
            "未宣称随机化已验证": bool(actual["limitations"])}, actual=actual)
    case("experiment_user_welch", "experiment", "重复事件按独立用户均值比较", experiment_units, {"control_mean": 3, "treatment_mean": 5, "difference": 2, "p": 1 - 1 / math.sqrt(2), "df": 2, "unit": "user_id"}, fixture="8 events -> user means [2,4] vs [4,6]", evidence=("样本量", "统计单位", "p值", "置信区间", "限制"))
    def experiment_direction():
        reverse = analyze_ab_test(experiment, "group", "outcome", statistical_unit="user_id", control_value="treatment", treatment_value="control")
        mixed = experiment.copy(); mixed.loc[4, "user_id"] = 1
        try: analyze_ab_test(mixed, "group", "outcome", statistical_unit="user_id"); rejected=False
        except ValueError: rejected=True
        return _observed({"反向差为负2": _close(reverse["absolute_lift"], -2), "跨组同用户拒绝": rejected}, reverse_difference=reverse["absolute_lift"], cross_group_rejected=rejected)
    case("experiment_direction_and_overlap", "experiment", "反转对照方向且检查跨组重复用户", experiment_direction, {"reverse_difference": -2, "cross_group_rejected": True}, fixture="experiment8")
    def invalidation():
        session = AnalysisSession(); calls=[]
        def key(filters, revision="1"):
            return analysis_cache_key(dataset_id="fixture", dataset_revision=revision, config={"metric": "revenue", "filters": filters}, catalog_version="model1", computation_version="0.1.0")
        a=key([]); b=key([{"column":"city","operator":"eq","value":"甲"}]); c=key([], "2")
        for k in (a,a,b): session.run(k, lambda: (calls.append(k) or _run_period()))
        return _observed({"同配置命中不重算": len(calls)==2, "过滤变化失效": a!=b, "revision变化失效": a!=c}, calls=len(calls))
    case("cache_filter_revision_invalidates", "invalidation", "筛选改变与数据修订改变必须使缓存失效", invalidation, {"calls_for_A_A_B": 2, "keys_differ": True})
    def immutable():
        thread=AnalysisThread(); config={"question":"支付金额", "column_mapping": _mapping(), "parameters": _parameters()}
        node=thread.add(result(), config, dataset_id="fixture", revision="1", schema={"daily":list(frame)})
        config["column_mapping"]["metric_columns"]=["orders"]
        copy=node.config; copy["question"]="modified"
        return _observed({"源配置变更不覆盖历史": node.config["column_mapping"]["metric_columns"]==["revenue"],
            "取回副本修改不覆盖历史": node.config["question"]=="支付金额", "历史不含DataFrame": isinstance(node.result_summary,str) and isinstance(node.result_ref,str)}, stored_metric=node.config["column_mapping"]["metric_columns"])
    case("history_config_immutable", "invalidation", "编辑新分支不改变原运行配置与摘要", immutable, {"stored_metric": "revenue", "large_result_copies": 0})
    def release():
        session=AnalysisSession(); actual=result();session.accept("a",actual)
        node=session.history.add(actual,{"column_mapping":_mapping()},dataset_id="fixture",revision="1",schema={"daily":list(frame)})
        before=session.history.result_available(node,session.current);session.clear()
        return _observed({"清理前可读原结果":before,"释放后不冒充结果存在":not session.history.result_available(node,session.current),
            "仍保留复用配置":node.config["column_mapping"]["metric_columns"]==["revenue"],"仅一大结果预算":session.cache.max_entries==1}, summary_rows=len(node.summary["rows"]), result_available=False)
    case("released_result_restore_config_only", "release", "释放旧大结果后只能恢复配置与摘要", release, {"result_available": False, "config_available": True})
    def old_task():
        thread=AnalysisThread(); thread.bind_dataset("fixture","1")
        old=thread.begin_request();thread.begin_request()
        try:thread.add(result(),{},dataset_id="fixture",revision="1",schema={},node_id=old);rejected=False
        except ValueError:rejected=True
        other=AnalysisThread()
        return _observed({"旧node不采纳":rejected,"其它会话不产生节点":not other.nodes,"当前历史仍空":not thread.nodes},rejected=rejected)
    case("stale_task_node_rejected", "release", "旧后台node不能覆盖新分支", old_task, {"rejected": True, "new_nodes": 0})
    def export_all():
        from insightpilot.reports.manifest import RunManifest
        from insightpilot.reports.pdf import generate_pdf_report
        from insightpilot.reports.markdown import generate_markdown_report
        from insightpilot.reports.html import generate_html_report
        from insightpilot.reports.excel import generate_excel_report
        from insightpilot.reports.bundle import generate_export_bundle
        from pypdf import PdfReader
        from openpyxl import load_workbook
        actual=result(); manifest=RunManifest.from_dict(actual["run_manifest"])
        pdf=generate_pdf_report(actual,manifest);md=generate_markdown_report(actual);html=generate_html_report(actual,[],manifest);xlsx=generate_excel_report(actual,manifest);js=manifest.to_json()
        bundle=generate_export_bundle(md,html,xlsx,js,actual["result_tables"],pdf_bytes=pdf)
        text="\n".join(p.extract_text() or "" for p in PdfReader(BytesIO(pdf)).pages)
        workbook=load_workbook(BytesIO(xlsx),read_only=True,data_only=False)
        excel_text=" ".join(str(v) for row in workbook["指标口径"].iter_rows(values_only=True) for v in row if v is not None); excel_run=dict(workbook["运行清单"].iter_rows(min_row=2,values_only=True)).get("run_id"); workbook.close()
        with ZipFile(BytesIO(bundle)) as archive:
            names=set(archive.namelist());inside=json.loads(archive.read("manifest.json"));same_pdf=archive.read("report.pdf")==pdf
        return _observed({"六格式有原run":all(manifest.run_id in content for content in (md,html,js,text)) and inside["run_id"]==manifest.run_id and excel_run==manifest.run_id,
            "口径进入所有文本报告":all("指标口径" in content and "revenue" in content for content in (md,html,text)) and "revenue" in excel_text and "口径版本" in excel_text and "元" in excel_text,
            "manifest保留原定义": bool(manifest.metric_definitions) and manifest.metric_definitions[0]["metric_id"]=="revenue",
            "ZIP含一致PDF且无源数据":same_pdf and {"report.pdf","report.md","report.html","report.xlsx","manifest.json"}.issubset(names) and not any("daily" in n or n.endswith(".db") for n in names)},
            run_id=manifest.run_id,bytes={"pdf":len(pdf),"xlsx":len(xlsx),"zip":len(bundle)},zip_members=sorted(names))
    case("exports_six_bound_to_run", "exports", "同一运行的六类导出保留原口径、运行编号与证据", export_all, {"formats": 6, "metric": "revenue", "run_binding": True, "raw_source": False}, evidence=("原运行编号", "口径版本", "证据", "ZIP目录"),forbidden=("当前编辑配置覆盖原口径", "泄露原始输入表"))
    def pdf_only():
        from insightpilot.reports.pdf import generate_pdf_report
        from insightpilot.reports.manifest import RunManifest
        actual=result();manifest=RunManifest.from_dict(actual["run_manifest"])
        with patch("insightpilot.reports.html.generate_html_report",side_effect=AssertionError("PDF调用HTML")) as html, patch("insightpilot.reports.excel.generate_excel_report",side_effect=AssertionError("PDF调用Excel")) as excel:
            content=generate_pdf_report(actual,manifest)
        wrong=replace(manifest,run_id="different-run")
        try:generate_pdf_report(actual,wrong);rejected=False
        except ValueError:rejected=True
        return _observed({"PDF独立生成":content.startswith(b"%PDF") and html.call_count==excel.call_count==0,"拒绝混合run":rejected},other_export_calls=html.call_count+excel.call_count)
    case("exports_pdf_only_and_wrong_run", "exports", "仅PDF不触发其它格式，拒绝混合不同运行", pdf_only, {"other_export_calls": 0, "mismatched_run_rejected": True})
    def invariants():
        shuffled=float(_daily_series(ratios.sample(frac=1,random_state=9),"ctr")["metric_value"].iloc[0])
        duplicated=float(_daily_series(pd.concat([people,people]),"active_users")["metric_value"].iloc[0])
        return _observed({"行顺序不改加权比率":_close(shuffled,81/110),"明确去重后重复主键不改人数":duplicated==3},weighted=shuffled,distinct=duplicated)
    case("metamorphic_shuffle_and_declared_dedup", "weighted", "打乱行序与明确去重下重复实体", invariants, {"ratio":81/110,"distinct":3},fixture="ratio2 and people5")
    def display_only():
        from insightpilot.metrics.definitions import computation_payload, stable_fingerprint
        card=result()["metric_definitions"][0]; changed={**card,"display_name":"展示标题","color":"blue"}
        same=stable_fingerprint(computation_payload(card))==stable_fingerprint(computation_payload(changed))
        session=AnalysisSession();calls=[]
        def actual_key(definition, semantic):
            return analysis_cache_key(dataset_id="fixture",dataset_revision="1",config={"metric_card":definition,"semantic_fingerprint":semantic},catalog_version="catalog1",computation_version="0.1.0")
        semantic=result()["semantic_fingerprint"]
        keys=[actual_key(card,semantic),actual_key(changed,semantic),actual_key(card,semantic)]
        for key in keys:session.run(key,lambda:(calls.append(1) or _run_period()))
        return _observed({"标题颜色不影响计算语义":same,"真实缓存键保持相同":len(set(keys))==1,"展示切换不重跑":len(calls)==1,"语义指纹变化失效":actual_key(card,"changed-semantic")!=keys[0]},analysis_calls=len(calls))
    case("metamorphic_display_no_recompute", "invalidation", "仅修改标题和颜色不改变数值或调用次数", display_only, {"analysis_calls":1,"semantic_fingerprint_unchanged":True})
    def alias_mapping():
        with AnalyticsEngine() as engine:
            engine.register_tables({"a": pd.DataFrame({"金额": [10, 20]}), "b": pd.DataFrame({"gross_amount": [10, 20]})})
            first=float(engine.run_sql('SELECT SUM("金额") AS total FROM a').iloc[0,0])
            second=float(engine.run_sql('SELECT SUM(gross_amount) AS total FROM b').iloc[0,0])
        return _observed({"安全别名映射不改数值":first==second==30},original=first,renamed=second)
    case("metamorphic_safe_column_alias", "users", "显式安全字段映射重命名后金额保持30", alias_mapping, {"total_before":30,"total_after":30},fixture="alias2:[10,20]")
    def binary_points():
        from insightpilot.ui.experiment_display import experiment_display
        binary=pd.DataFrame({"user_id":list(range(8)),"group":["control"]*4+["treatment"]*4,"completion_rate":[0,0,0,1,0,1,1,1]})
        actual=run_agent_analysis("按用户比较二元完播率",{"experiment":binary},data_source_type="csv",goal_mode="experiment_analysis",playbook_id="experiment_comparison",column_mapping={"table_name":"experiment","metric_columns":["completion_rate"],"group_column":"group","statistical_unit":"user_id"},playbook_parameters={"metric":"completion_rate","metric_type":"proportion","control_value":"control","treatment_value":"treatment","confidence_level":.95},presentation_mode="deferred")
        item=actual["analysis_result_package"]["experiment_results"][0]
        display=experiment_display(item)
        se=math.sqrt(3/32); expected_p=float(2*stats.norm.sf(.5/se)); half=float(stats.norm.ppf(.975))*se
        return _observed({"真实实验完成":actual["execution_status"]=="COMPLETED","独立用户4比4":item["control_sample_size"]==item["treatment_sample_size"]==4,"绝对差0.5":_close(item["absolute_lift"],.5),"显示50个百分点":display["绝对提升（lift）"]=="50 个百分点","同Wald口径p及CI":_close(item["p_value"],expected_p) and _close(item["confidence_interval_lower"],.5-half) and _close(item["confidence_interval_upper"],.5+half)},actual=item,display=display)
    case("experiment_binary_percentage_points", "experiment", "二元用户完播从25%到75%，差为50个百分点", binary_points, {"control":.25,"treatment":.75,"difference":.5,"display_points":50,"unit":"user_id","n_each":4},fixture="binary8:[0,0,0,1] vs [0,1,1,1]",evidence=("独立样本量","绝对差百分点","p值与区间同口径"))
    return cases


def run_task_evaluation() -> EvaluationSuiteResult:
    from insightpilot.evaluation.guided_suite import build_guided_task_cases
    result=EvaluationHarness().run(build_task_cases() + build_guided_task_cases())
    return EvaluationSuiteResult(result.suite,result.overall_score,result.results,{"task_contract":1.0})
