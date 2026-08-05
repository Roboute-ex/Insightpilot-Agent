from __future__ import annotations

from io import BytesIO
from zipfile import ZipFile

import openpyxl

from insightpilot.agents.workflow import run_agent_analysis
from insightpilot.data.synthetic import generate_multi_table_commerce_data
from insightpilot.reports.bundle import generate_export_bundle
from insightpilot.reports.excel import generate_excel_report
from insightpilot.reports.html import generate_html_report
from insightpilot.reports.manifest import RunManifest


def _semantic_result():
    return run_agent_analysis(
        "按城市分析总金额",
        generate_multi_table_commerce_data(seed=42),
        playbook_id="semantic_metric_query",
        playbook_parameters={"metric": "total_revenue", "dimensions": ["customer_city"]},
    )


def test_markdown_and_html_reports_are_chinese_first() -> None:
    result = _semantic_result()
    report = result["report_markdown"]
    for heading in ["分析问题", "数据来源", "分析目标", "分析方案", "语义指标", "多表连接方案", "核心发现", "可视化诊断", "分析质量检查", "数据质量与 Contract", "风险与限制", "数据血缘", "运行摘要", "可复现配置", "后续建议"]:
        assert heading in report
    manifest = RunManifest.from_dict(result["run_manifest"])
    html = generate_html_report(result, result["charts"], manifest)
    assert 'lang="zh-CN"' in html
    assert "中文分析报告" in html
    assert "多表连接方案" in html


def test_excel_and_zip_readme_are_chinese_first_with_compatibility_aliases() -> None:
    result = _semantic_result()
    manifest = RunManifest.from_dict(result["run_manifest"])
    excel = generate_excel_report(result, manifest)
    workbook = openpyxl.load_workbook(BytesIO(excel), read_only=True)
    assert {"分析摘要", "核心发现", "质量检查", "字段映射", "查询计划", "数据血缘", "运行清单"}.issubset(workbook.sheetnames)
    assert {"Summary", "Findings", "Reviewer", "Mapping", "Manifest"}.issubset(workbook.sheetnames)
    bundle = generate_export_bundle(result["report_markdown"], "<html></html>", excel, manifest.to_json(), result["result_tables"])
    with ZipFile(BytesIO(bundle)) as archive:
        readme = archive.read("README.txt").decode("utf-8")
    assert "中文分析导出包" in readme
    assert "不包含上传源文件" in readme
