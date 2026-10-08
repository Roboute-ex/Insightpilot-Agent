from __future__ import annotations

from io import BytesIO

import openpyxl

from insightpilot.agents.workflow import run_agent_analysis
from insightpilot.reports.excel import generate_excel_report
from insightpilot.reports.manifest import RunManifest


def test_excel_export_can_be_reopened(playbook_tables, playbook_mapping) -> None:
    result = run_agent_analysis("profile", playbook_tables, column_mapping=playbook_mapping, data_source_type="uploaded_files", playbook_id="data_profile")
    content = generate_excel_report(result, RunManifest.from_dict(result["run_manifest"]), compatibility_mode=True)
    workbook = openpyxl.load_workbook(BytesIO(content), read_only=True)
    assert {"Summary", "Findings", "Reviewer", "Mapping", "Manifest"}.issubset(workbook.sheetnames)
    assert any(name.startswith("Result_") for name in workbook.sheetnames)
