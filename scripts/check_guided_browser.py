"""Real isolated Chrome guided-analysis acceptance; independent of AppTest adapters."""
from __future__ import annotations
import argparse,json,time,urllib.request
from pathlib import Path
from playwright.sync_api import sync_playwright
from streamlit.proto.ForwardMsg_pb2 import ForwardMsg

def main():
 p=argparse.ArgumentParser();p.add_argument('--url',required=True);p.add_argument('--output',type=Path,required=True);p.add_argument('--before',action='store_true');args=p.parse_args()
 args.output.mkdir(parents=True,exist_ok=False)
 record={'url':args.url,'before':args.before,'health_status':urllib.request.urlopen(args.url+'/_stcore/health').status,'checks':{},'operations':[]}
 with sync_playwright() as pw:
  browser=pw.chromium.launch(executable_path=r'C:\Program Files\Google\Chrome\Application\chrome.exe',headless=True,chromium_sandbox=True)
  record['chrome']=browser.version;context=browser.new_context(viewport={'width':1500,'height':1060},accept_downloads=True);page=context.new_page();errors=[];finished=[0]
  page.on('pageerror',lambda error:errors.append(str(error)))
  def frame(raw):
   if isinstance(raw,bytes):
    message=ForwardMsg()
    try:message.ParseFromString(raw)
    except Exception:return
    if message.WhichOneof('type')=='script_finished' and message.script_finished in (0,3):finished[0]+=1
  page.on('websocket',lambda ws:ws.on('framereceived',frame))
  def idle():
   page.locator('[data-testid="stStatusWidget"]').wait_for(state='hidden',timeout=90000)
   page.evaluate('()=>new Promise(r=>requestAnimationFrame(()=>requestAnimationFrame(r)))')
  def action(name,fn,wait=True):
   idle();old=finished[0];start=time.perf_counter();fn()
   if wait:
    end=time.perf_counter()+90
    while finished[0]<=old:
     if time.perf_counter()>end:raise TimeoutError(name+' did not finish a Streamlit rerun')
     page.wait_for_timeout(20)
   idle();record['operations'].append({'name':name,'seconds':time.perf_counter()-start});print(name,flush=True)
  def click(label):action(label,lambda:page.get_by_role('button',name=label,exact=False).click())
  def expand(label):action(label,lambda:page.get_by_text(label,exact=True).click())
  def choose(label,value):
   def operate():
    loc=page.get_by_role('combobox',name=label,exact=True);loc.fill(value);loc.press('ArrowDown');page.get_by_role('option',name=value,exact=True).click()
   action('select '+label+' '+value,operate)
  def capture(name):page.screenshot(path=str(args.output/(name+'.png')),full_page=True)
  try:
   page.goto(args.url);page.get_by_role('button',name='开始分析').wait_for(timeout=90000);idle();capture('fresh')
   record['checks']['no_initial_result']=not page.get_by_text('执行摘要',exact=True).count()
   if args.before:
    record['checks']['fresh_default']=page.get_by_role('combobox',name='分析剧本',exact=True).input_value()
    action('set question',lambda:(page.get_by_role('textbox',name='分析问题',exact=True).fill('昨日订单量为什么下降？'),page.get_by_role('textbox',name='分析问题',exact=True).press('Enter')))
    choose('分析目标模式','指标波动诊断');choose('分析剧本','轻量因果探索')
    record['checks']['selected_causal']=page.get_by_role('combobox',name='分析剧本',exact=True).input_value()
    capture('causal-config');click('开始分析');page.get_by_text('当前任务：失败',exact=True).wait_for(timeout=90000);capture('blocked')
    record['checks']['generic_failed']=True
    click('填入此问题');click('开始分析');page.get_by_text('分析已完成。以下内容由当前数据实际计算生成。',exact=True).wait_for(timeout=90000)
   else:
    record['checks']['no_modes']=not page.get_by_text('简洁演示',exact=True).count() and not page.get_by_text('专业分析',exact=True).count()
    record['checks']['one_start_button']=page.get_by_role('button',name='开始分析').count()==1
    record['checks']['no_initial_tables']=page.locator('[data-testid="stDataFrame"]').count()==0
    record['checks']['no_initial_groups']=not page.get_by_text('处理组值',exact=True).count()
    question=page.get_by_role('textbox',name='你想分析什么？',exact=True)
    action('set question',lambda:(question.fill('昨日订单量为什么下降？'),question.press('Enter')))
    expand('调整方法与参数');choose('分析方法','轻量因果探索')
    page.get_by_text('treatment、control 只是组别取值，不能代替数据表中的分组字段。',exact=False).first.wait_for()
    record['checks']['blocked_disabled']=not page.get_by_role('button',name='开始分析').is_enabled()
    record['checks']['no_false_missing_outcome']='缺少结果变量' not in page.locator('body').inner_text()
    record['checks']['no_failed_task']='当前任务：失败' not in page.locator('body').inner_text();capture('causal-settings')
    expand('调整方法与参数');record['checks']['manual_visible_when_closed']='已手动选择：轻量因果探索' in page.locator('body').inner_text()
    page.get_by_role('button',name='改用推荐方法').scroll_into_view_if_needed();capture('blocked')
    click('改用推荐方法');record['checks']['accept_did_not_run']=not page.get_by_text('执行摘要',exact=True).count();capture('accepted-ready')
    click('开始分析');page.get_by_text('执行摘要',exact=True).wait_for(timeout=90000)
   idle();capture('completed');record['checks']['complete']=page.get_by_text('执行摘要',exact=True).is_visible()
   if not args.before:
    expand('查看口径与质量');capture('quality')
    record['font_samples']={key:page.locator(selector).first.evaluate('(e)=>getComputedStyle(e).fontSize') for key,selector in [('title','.ip-summary-title'),('body','.ip-summary-content')]}
    record['checks']['original_font_sizes']=record['font_samples']=={'title':'16px','body':'14px'}
    expand('查看口径与质量')
    # Navigation is native segmented control; no session-state injection.
    action('report navigation',lambda:page.get_by_text('报告',exact=True).click())
    record['checks']['lazy_exports']=page.get_by_role('button',name='下载 PDF 报告').count()==0
    for label,kind in [('PDF 报告','pdf'),('Excel 报告','excel'),('ZIP 报告包','zip')]:
     click('生成 '+label)
     with page.expect_download() as download:page.get_by_role('button',name='下载 '+label).click()
     path=args.output/('selected-run.'+{'excel':'xlsx'}.get(kind,kind));download.value.save_as(path);record['checks']['download_'+kind]=path.stat().st_size>0
    import zipfile
    from pypdf import PdfReader
    from openpyxl import load_workbook
    with zipfile.ZipFile(args.output/'selected-run.zip') as bundle:
     pdf_name=next(name for name in bundle.namelist() if name.endswith('report.pdf'))
     record['checks']['zip_pdf_same_bytes']=bundle.read(pdf_name)==(args.output/'selected-run.pdf').read_bytes()
    record['checks']['readable_pdf']=len(PdfReader(args.output/'selected-run.pdf').pages)>0
    workbook=load_workbook(args.output/'selected-run.xlsx',read_only=True)
    record['checks']['readable_excel']=bool(workbook.sheetnames);workbook.close()
    capture('exports')
   record['page_errors']=errors;assert not errors
   assert all(v for v in record['checks'].values()),record['checks']
  except Exception as exc:
   record['error']=str(exc);record['body']=page.locator('body').inner_text();capture('failure');raise
  finally:
   (args.output/'checks.json').write_text(json.dumps(record,ensure_ascii=False,indent=2)+'\n',encoding='utf-8');context.close();browser.close()
 print(json.dumps(record['checks'],ensure_ascii=False));return 0
if __name__=='__main__':raise SystemExit(main())
