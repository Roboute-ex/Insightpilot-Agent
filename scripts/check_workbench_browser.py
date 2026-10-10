"""Real Chrome workbench acceptance: native controls/uploads, no state/result injection.

Run against a separately started local Streamlit server. Each output directory is
new. Screenshots, downloaded synthetic reports and failures remain local reports.
"""
from __future__ import annotations
import argparse
import csv
import json
import re
from pathlib import Path
import time
import urllib.request
import zipfile

from playwright.sync_api import sync_playwright
from streamlit.proto.ForwardMsg_pb2 import ForwardMsg


class BrowserProbe:
    def __init__(self, browser, output, record):
        self.browser, self.output, self.record = browser, output, record
        self.contexts = []; self.page = None; self.finished = 0; self.errors = []

    def session(self, url):
        context = self.browser.new_context(viewport={"width": 1440, "height": 900}, accept_downloads=True)
        self.contexts.append(context); self.page = context.new_page(); self.finished = 0
        self.page.on("pageerror", lambda error: self.errors.append(str(error)))
        def frame(raw):
            if not isinstance(raw, bytes): return
            message = ForwardMsg()
            try: message.ParseFromString(raw)
            except Exception: return
            if message.WhichOneof("type") == "script_finished" and message.script_finished in (0, 3):
                self.finished += 1
        self.page.on("websocket", lambda ws: ws.on("framereceived", frame))
        self.page.goto(url, wait_until="domcontentloaded")
        self.page.get_by_role("button", name="开始分析", exact=False).wait_for(timeout=90000)
        self.idle()

    def idle(self):
        self.page.locator('[data-testid="stStatusWidget"]').wait_for(state="hidden", timeout=90000)
        self.page.evaluate("()=>new Promise(r=>requestAnimationFrame(()=>requestAnimationFrame(r)))")

    def action(self, name, fn, *, wait=True):
        self.idle(); old = self.finished; started = time.perf_counter(); fn()
        if wait:
            deadline = time.perf_counter() + 90
            while self.finished <= old:
                if time.perf_counter() > deadline: raise TimeoutError(name + " did not finish a Streamlit rerun")
                self.page.wait_for_timeout(20)
        self.idle()
        self.record["operations"].append({"name": name, "seconds": time.perf_counter()-started})
        print(name, flush=True)

    def click(self, label):
        self.action(label, lambda: self.page.get_by_role("button", name=label, exact=False).first.click())

    def navigate(self, label):
        target=self.page.get_by_text(label,exact=True).first
        if target.evaluate('(e)=>e.closest("button")?.getAttribute("aria-checked")==="true"'):
            return
        self.action("导航："+label,lambda:target.click())

    def choose(self, label, value, *, wait=True):
        control = self.page.get_by_role("combobox", name=label, exact=True)
        if control.input_value() == value: return
        def operate():
            control.click();control.press_sequentially(value)
            self.page.get_by_role("option",name=value,exact=True).click()
        self.action("选择："+label+"="+value, operate, wait=wait)

    def multiselect(self, label, value, *, wait=True):
        control=self.page.get_by_role("combobox",name=label,exact=True)
        widget=control.locator('xpath=ancestor::*[@data-testid="stMultiSelect"][1]')
        if widget.get_by_text(value,exact=True).count():
            self.record.setdefault("existing_multi_selections",[]).append({"label":label,"value":value})
            return
        def operate():
            control.click();control.press_sequentially(value)
            self.page.get_by_role("option",name=value,exact=True).click()
        self.action("增加："+label+"="+value,operate,wait=wait)
        self.page.keyboard.press("Escape");self.idle()
        # The fixture explicitly requests one field. Remove any visible extra
        # chip through its real button, never by mutating application state.
        names=widget.locator("[data-tag]").evaluate_all('(tags)=>tags.map(t=>t.getAttribute("aria-label"))')
        for name in names:
            if name and name!=value:
                self.action("移除："+label+"="+name,lambda name=name:widget.get_by_role("button",name="Remove "+name,exact=True).click(),wait=wait)
        selected=widget.locator("[data-tag]").evaluate_all('(tags)=>tags.map(t=>t.getAttribute("aria-label"))')
        self.record.setdefault("actual_multi_selections",[]).append({"label":label,"values":selected})
        self.check("configured_"+label+"_single_"+value,selected==[value])

    def question(self, value):
        control = self.page.get_by_role("textbox", name="你想分析什么？", exact=True)
        self.action("修改问题", lambda: (control.fill(value), control.press("Enter")))

    def body(self): return self.page.locator("body").inner_text()

    def check(self, name, condition):
        self.record["checks"][name] = bool(condition)
        if not condition: raise AssertionError(name)

    def capture(self, name):
        self.idle()
        self.page.screenshot(path=str(self.output/(name+".png")), full_page=True)
        geometry = self.page.evaluate("""()=>({width:innerWidth,height:innerHeight,dpr:devicePixelRatio,
            visualScale:visualViewport.scale,documentWidth:document.documentElement.scrollWidth,
            overflow:document.documentElement.scrollWidth>innerWidth,
            visibleBodyFontSizes:[...new Set([...document.querySelectorAll('p')]
              .filter(x=>x.getBoundingClientRect().width&&x.getBoundingClientRect().height)
              .map(x=>getComputedStyle(x).fontSize))]})""")
        self.record.setdefault("viewports", {})[name] = geometry
        self.check(name+"_100percent", geometry["visualScale"] == 1)
        self.check(name+"_no_page_overflow", not geometry["overflow"])

    def viewport_checks(self, prefix):
        for width, height in [(1366,768),(1440,900),(1920,1080),(1024,768)]:
            self.page.set_viewport_size({"width": width, "height": height})
            self.capture(f"{prefix}-{width}x{height}")
        self.page.set_viewport_size({"width":1440,"height":900})

    def export(self, prefix, *, all_formats=True):
        self.navigate("报告")
        generators = [text.strip() for text in self.page.get_by_role("button").all_text_contents() if text.strip().startswith("生成") and text.strip() not in {"生成探索结果", "生成分布"}]
        expected = ["生成 PDF 报告", "生成 Markdown 报告", "生成 HTML 报告", "生成 Excel 报告", "生成 运行清单 JSON", "生成 ZIP 报告包"]
        self.check(prefix+"_six_pdf_first", generators == expected)
        self.check(prefix+"_lazy_pdf", self.page.get_by_role("button", name="下载 PDF 报告").count() == 0)
        selected = [("PDF 报告","pdf"),("Markdown 报告","md"),("HTML 报告","html"),("Excel 报告","xlsx"),("运行清单 JSON","json"),("ZIP 报告包","zip")]
        if not all_formats: selected = [item for item in selected if item[1] in {"pdf","xlsx","json","zip"}]
        paths = {}
        for label, extension in selected:
            self.click("生成 "+label)
            with self.page.expect_download() as download:
                self.page.get_by_role("button", name="下载 "+label, exact=False).click()
            target = self.output/(prefix+"."+extension); download.value.save_as(target)
            self.check(prefix+"_download_"+extension, target.stat().st_size > 0)
            paths[extension] = target
        manifest = json.loads(paths["json"].read_text(encoding="utf-8"))
        from pypdf import PdfReader
        pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(paths["pdf"]).pages)
        self.check(prefix+"_run_bound_pdf", manifest["run_id"] in pdf_text)
        with zipfile.ZipFile(paths["zip"]) as bundle:
            pdf_name = next(name for name in bundle.namelist() if name.endswith("report.pdf"))
            self.check(prefix+"_zip_pdf_exact", bundle.read(pdf_name) == paths["pdf"].read_bytes())
        self.record.setdefault("exports", {})[prefix] = {"run_id": manifest["run_id"], "paths": {k:str(v) for k,v in paths.items()}, "pdf_text":pdf_text, "manifest":manifest}
        return paths, manifest, pdf_text

    def close(self):
        for context in self.contexts: context.close()


def causal_fixture(path):
    with path.open("w",newline="",encoding="utf-8") as stream:
        writer=csv.DictWriter(stream,fieldnames=["date","group","outcome","covariate"]);writer.writeheader()
        for group in ["control","treatment"]:
            for x in range(12):
                writer.writerow({"date":f"2026-09-{x+1:02d}","group":group,"outcome":10+x+(2 if group=="treatment" else 0)+(0.25 if x%2 else -0.25),"covariate":x})


def catalog_and_order(probe):
    p=probe
    p.check("no_legacy_modes", not p.page.get_by_text("简洁演示",exact=True).count())
    p.check("four_navigation_labels", all(p.page.get_by_text(label,exact=True).count() for label in ["分析工作台","数据探索","分析方法","历史比较"]))
    p.viewport_checks("01-home")
    p.question("昨日订单量为什么下降？")
    p.navigate("分析方法")
    labels=p.page.get_by_role("button",name="配置：").all_text_contents()
    p.record["catalog_buttons"]=labels;p.check("catalog_all_ten",len(labels)==10)
    p.viewport_checks("02-methods")
    p.choose("方法分类","实验与效果")
    p.check("category_two_effect_methods",len(p.page.get_by_role("button",name="配置：").all_text_contents())==2)
    search=p.page.get_by_role("textbox",name="搜索分析方法",exact=True)
    p.action("搜索因果",lambda:(search.fill("因果"),search.press("Enter")))
    p.check("search_one_causal",len(p.page.get_by_role("button",name="配置：").all_text_contents())==1)
    p.click("清除方法筛选");p.check("clear_returns_ten",len(p.page.get_by_role("button",name="配置：").all_text_contents())==10)
    p.click("配置：轻量因果探索")
    p.check("configuration_keeps_orders",p.page.get_by_role("combobox",name="结果字段",exact=True).input_value()=="orders")
    p.check("configuration_keeps_daily_metrics",p.page.get_by_role("combobox",name="分析范围：目标数据表",exact=True).input_value()=="daily_metrics")
    p.check("invalid_causal_disabled",not p.page.get_by_role("button",name="开始分析",exact=False).is_enabled())
    p.check("group_values_not_fields", "treatment、control 只是组别取值" in p.body())
    p.choose("结果字段","不选择")
    p.check("explicit_empty_outcome_blocked",not p.page.get_by_role("button",name="开始分析",exact=False).is_enabled() and "缺少结果" in p.body())
    p.page.get_by_role("button",name="开始分析",exact=False).scroll_into_view_if_needed();p.capture("03-invalid-causal")
    p.choose("结果字段","orders")
    p.click("改用推荐方法")
    p.check("accept_recommendation_does_not_execute", not p.page.get_by_text("执行摘要",exact=True).count())
    p.click("开始分析");p.page.get_by_text("执行摘要",exact=True).wait_for(timeout=90000);p.idle()
    p.check("order_completed", p.page.get_by_text("已完成", exact=True).count() >= 1 and all(term in p.body() for term in ["执行摘要", "昨日订单量为", "核心指标表现"]))
    p.page.get_by_text("执行摘要",exact=True).scroll_into_view_if_needed();p.viewport_checks("04-order-result")


def valid_causal(probe,url,fixture):
    p=probe;p.session(url);p.choose("数据来源","上传 CSV / Excel")
    p.action("上传24行独立合成因果CSV",lambda:p.page.locator('input[type="file"]').set_input_files(str(fixture)))
    p.page.get_by_role("textbox",name="你想分析什么？",exact=True).wait_for(timeout=90000);p.idle()
    p.question("控制协变量后，处理组与对照组的结果有多大差异？")
    p.navigate("分析方法");p.click("配置：轻量因果探索");p.choose("分析目标","轻量因果探索")
    p.choose("分组字段","不选择");p.choose("处理字段","group");p.choose("结果字段","outcome");p.choose("聚合方式","均值")
    p.multiselect("指标字段","outcome");p.multiselect("协变量","covariate");p.click("确认当前字段映射")
    p.multiselect("控制变量","covariate",wait=False);p.choose("处理组值","treatment",wait=False);p.choose("对照组值","control",wait=False)
    p.click("应用参数并更新建议")
    p.check("causal_ready",p.page.get_by_role("button",name="开始分析",exact=False).is_enabled())
    p.page.get_by_role("button",name="开始分析",exact=False).scroll_into_view_if_needed();p.viewport_checks("06-causal-ready")
    p.click("开始分析");p.page.get_by_text("执行摘要",exact=True).wait_for(timeout=90000);p.idle()
    p.check("causal_completed", p.page.get_by_text("已完成", exact=True).count() >= 1 and p.page.get_by_text("执行摘要", exact=True).count() == 1)
    p.record["causal_result_body"]=p.body()
    p.check("causal_summary_actual",all(term in p.body() for term in ["=2.0000","n=24","此结果不是因果证明"]))
    p.page.get_by_text("执行摘要",exact=True).scroll_into_view_if_needed();p.capture("07-causal-result")
    p.page.locator('[data-testid="stPlotlyChart"]').first.scroll_into_view_if_needed();p.capture("07b-causal-grouped-chart")
    paths,manifest,pdf_text=p.export("causal")
    from openpyxl import load_workbook
    workbook=load_workbook(paths["xlsx"],read_only=True,data_only=True)
    rows=list(workbook["结果_adjusted_effect_summary"].values);workbook.close()
    actual=dict(zip(rows[0],rows[1]));p.record["causal_actual"]=actual
    p.check("causal_naive_2",abs(float(actual["naive_difference"])-2)<1e-10)
    p.check("causal_adjusted_2",abs(float(actual["adjusted_effect"])-2)<1e-10)
    p.check("causal_n_24",actual["样本量"]==24)
    p.check("causal_actual_covariate",actual["covariates"]=="['covariate']")
    p.check("causal_pdf_limitations",all(term in pdf_text for term in ["无未观测混杂","重叠性","此结果不是因果证明"]))
    p.capture("08-causal-exports")


def exploration_fixture(path):
    with path.open("w",newline="",encoding="utf-8") as stream:
        writer=csv.DictWriter(stream,fieldnames=["row_group","column_group","amount","entity_id","numerator","denominator"]);writer.writeheader()
        for row in [
            {"row_group":"甲","column_group":"X","amount":10,"entity_id":"u1","numerator":9,"denominator":90},
            {"row_group":"甲","column_group":"Y","amount":30,"entity_id":"u2","numerator":1,"denominator":2},
            {"row_group":"乙","column_group":"X","amount":20,"entity_id":"u1","numerator":0,"denominator":0},
            {"row_group":"乙","column_group":"Y","amount":60,"entity_id":"u3","numerator":0,"denominator":0},
        ]: writer.writerow(row)


def exploration_run_text(probe):
    locator=probe.page.get_by_text(re.compile(r"^运行 [0-9a-f-]+；"))
    return locator.first.inner_text() if locator.count() else None


def wait_new_exploration(probe, old=None):
    deadline=time.perf_counter()+90
    while True:
        probe.idle()
        current=exploration_run_text(probe)
        if current and current!=old: return current
        if time.perf_counter()>deadline: raise TimeoutError("No new completed exploration run: "+probe.body()[-2500:])
        probe.page.wait_for_timeout(100)


def pivot_exploration(probe,url,fixture):
    p=probe;p.session(url);p.choose("数据来源","上传 CSV / Excel")
    p.action("上传四行独立手算透视CSV",lambda:p.page.locator('input[type="file"]').set_input_files(str(fixture)))
    p.page.get_by_role("textbox",name="你想分析什么？",exact=True).wait_for(timeout=90000);p.idle()
    p.navigate("数据探索");p.check("fields_visible","字段概览" in p.body());p.capture("09-fields")
    p.navigate("交叉表")
    p.choose("指标字段","amount",wait=False)
    p.choose("行维度 / 分组字段","row_group",wait=False)
    p.choose("列维度","column_group",wait=False)
    p.action("确认探索口径",lambda:p.page.get_by_role("checkbox",name="我确认以上字段、聚合口径与范围，用于本次描述性探索",exact=True).press("Space"),wait=False)
    p.click("生成探索结果");wait_new_exploration(p)
    p.page.get_by_text("已生成探索结果",exact=True).scroll_into_view_if_needed();p.viewport_checks("10-pivot-sum")
    p.check("pivot_sum_summary","120" in p.body())
    paths,manifest,pdf_text=p.export("pivot-sum")
    verify_pivot_excel(p,paths["xlsx"],120,{"甲":40,"乙":80},"sum")
    # The same four rows are recomputed only after an explicit new request.
    for aggregation,label,metric,expected,groups in [
        ("mean","均值（按有效观测行）","amount",30,{"甲":20,"乙":40}),
        ("ratio","加权比率（总分子 / 总分母）","numerator",10/92,{"甲":10/92,"乙":None}),
        ("distinct","去重计数","entity_id",3,{"甲":2,"乙":2}),
    ]:
        old=exploration_run_text(p);p.choose("指标聚合口径",label)
        p.choose("指标字段",metric,wait=False)
        if aggregation=="ratio":
            p.choose("比率分子字段","numerator",wait=False);p.choose("比率分母字段","denominator",wait=False)
        p.action("重新确认探索口径",lambda:p.page.get_by_role("checkbox",name="我确认以上字段、聚合口径与范围，用于本次描述性探索",exact=True).press("Space"),wait=False)
        p.click("生成探索结果");wait_new_exploration(p,old)
        paths,manifest,pdf_text=p.export("pivot-"+aggregation)
        verify_pivot_excel(p,paths["xlsx"],expected,groups,aggregation)
    p.capture("11-pivot-exports")


def verify_pivot_excel(probe,path,expected,group_expected,name):
    from openpyxl import load_workbook
    workbook=load_workbook(path,read_only=True,data_only=True)
    target=next(sheet for sheet in workbook if "exploration_totals" in sheet.title)
    rows=list(target.values);records=[dict(zip(rows[0],row)) for row in rows[1:]];workbook.close()
    total=next(row for row in records if row["scope"]=="all")
    probe.check("pivot_"+name+"_independent_total",abs(float(total["metric_value"])-expected)<1e-10)
    for key,value in group_expected.items():
        actual=next(row["metric_value"] for row in records if row["scope"]=="row" and row["row_value"]==key)
        probe.check("pivot_"+name+"_group_"+key,actual is None if value is None else abs(float(actual)-value)<1e-10)
    probe.record.setdefault("pivot_actual",{})[name]=records


def drilldown(probe,parent_manifest):
    p=probe;p.navigate("结论")
    heading=p.page.get_by_text("按城市继续探索",exact=True)
    heading.wait_for(timeout=15000);heading.scroll_into_view_if_needed()
    # Streamlit's native single-row dataframe supports pointer selection. The
    # bounding box comes from the actual rendered selection table, never a forged event.
    table=p.page.locator('[class*="st-key-drilldown_select_"]').locator('[data-testid="stDataFrame"]').first
    if not table.count(): table=p.page.locator('[data-testid="stDataFrame"]').last
    table.scroll_into_view_if_needed();canvas=table.locator("canvas").first;box=canvas.bounding_box()
    if not box: raise RuntimeError("City selection table canvas is not rendered")
    p.action("选择城市分组（仅草稿）",lambda:p.page.mouse.click(box["x"]+16,box["y"]+54))
    p.check("selection_draft_only","筛选草稿：city = " in p.body() and "尚未执行" in p.body())
    # Run identity is intentionally hidden until technical details are opened.
    p.action("展开技术详情核对父运行", lambda: p.page.get_by_text("技术详情", exact=True).click())
    p.check("selection_keeps_parent", "运行编号："+parent_manifest["run_id"] in p.body())
    p.action("收起技术详情", lambda: p.page.get_by_text("技术详情", exact=True).click())
    p.capture("05-city-selection")
    p.click("基于选中分组继续分析");wait_new_exploration(p)
    p.check("drilldown_scope_caption","恢复运行的原筛选" in p.body())
    p.page.get_by_text("已生成探索结果",exact=True).scroll_into_view_if_needed();p.capture("05b-city-child")
    paths,child,pdf_text=p.export("city-child")
    p.check("child_new_run",child["run_id"]!=parent_manifest["run_id"])
    p.record["drilldown_parent_run_id"]=parent_manifest["run_id"]
    p.record["drilldown_child_manifest"]=child
    p.navigate("历史比较")
    p.check("history_two_runs",p.page.get_by_role("combobox",name="查看历史运行",exact=True).count()==1)
    p.click("返回上一步")
    p.check("history_selected_parent",parent_manifest["run_id"] in p.body())
    p.check("parent_released_visible","已释放" in p.body() or "需要重新执行" in p.body())
    p.capture("05c-history-released")


def main():
    parser=argparse.ArgumentParser();parser.add_argument("--url",required=True);parser.add_argument("--output",type=Path,required=True);parser.add_argument("--scenario",choices=["all","causal","pivot"],default="all")
    args=parser.parse_args();args.output.mkdir(parents=True,exist_ok=False)
    fixture=args.output/"causal_fixture.csv";causal_fixture(fixture)
    pivot_fixture=args.output/"pivot_fixture.csv";exploration_fixture(pivot_fixture)
    record={"scenario":args.scenario,"url":args.url,"health_status":urllib.request.urlopen(args.url+"/_stcore/health").status,"checks":{},"operations":[],"fixture_expected":{"rows":24,"naive_difference":2,"adjusted_effect":2,"basis":"paired groups share x=0..11 and alternating +/-0.25; outcome adds exactly2 for treatment"}}
    with sync_playwright() as playwright:
        browser=playwright.chromium.launch(executable_path=r"C:\Program Files\Google\Chrome\Application\chrome.exe",headless=True,chromium_sandbox=True)
        record["chrome"]=browser.version;probe=BrowserProbe(browser,args.output,record)
        try:
            if args.scenario=="all":
                probe.session(args.url);catalog_and_order(probe)
                _,parent_manifest,_=probe.export("order")
                drilldown(probe,parent_manifest)
            if args.scenario in {"all","causal"}: valid_causal(probe,args.url,fixture)
            if args.scenario in {"all","pivot"}: pivot_exploration(probe,args.url,pivot_fixture)
            record["page_errors"]=probe.errors;probe.check("no_page_errors",not probe.errors)
        except Exception as exc:
            record["error"]=str(exc)
            if probe.page is not None:
                record["failure_body"]=probe.body()
                probe.page.screenshot(path=str(args.output/"failure.png"),full_page=True)
            raise
        finally:
            (args.output/"checks.json").write_text(json.dumps(record,ensure_ascii=False,indent=2,default=str)+"\n",encoding="utf-8")
            probe.close();browser.close()
    print(json.dumps(record["checks"],ensure_ascii=False));return 0

if __name__=="__main__":raise SystemExit(main())
